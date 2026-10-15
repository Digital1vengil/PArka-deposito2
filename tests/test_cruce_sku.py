import io
import re

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app


def _login(c, nombre="Vengil", pin="1234"):
    uid = re.search(r"elegir\((\d+), '" + nombre + r"'\)", c.get("/login").text).group(1)
    assert c.post("/login", data={"usuario_id": uid, "pin": pin}).json()["ok"]


def _venta(db, sku, producto="", variante=""):
    from app.db import hoy
    from app.models import Control, ControlItem, ControlPaquete
    c = Control(fecha=hoy(), modo="ecommerce", tipo="flex", nombre="test", origen="api")
    db.add(c)
    db.flush()
    p = ControlPaquete(control_id=c.id, tracking=f"T{sku}", tipo="flex")
    db.add(p)
    db.flush()
    db.add(ControlItem(paquete_id=p.id, sku=sku, producto=producto, variante=variante))
    db.commit()


def test_cruce_sku():
    from app.db import SessionLocal
    from app.models import Articulo
    with SessionLocal() as db:
        db.add_all([
            Articulo(sku="1-CAMP-M-00777-BLCK-XL-PARK", codigo_barras="2999000000017", articulo="ZETA 00777",
                     variante="BLACK, XL, Parka", color="BLCK", talle="XL", odoo_id=880001),
            Articulo(sku="1-CAMP-M-00777-BLCK-XL-NOBR", codigo_barras="2999000000024", articulo="ZETA 00777",
                     variante="BLACK, XL, no brand", color="BLCK", talle="XL", odoo_id=880002),
            Articulo(sku="1-BUZO-M-ORIGZ-BLCK-L-PARK", codigo_barras="2999000000031", articulo="DUNE ORIGZ",
                     variante="L, Parka", color="BLCK", talle="L", odoo_id=880003),
        ])
        db.commit()
        _venta(db, "M-777-BLACK-XL", "Campera Zeta 00777 Parka")
        _venta(db, "BUZO-VIEJO-1", "Buzo Dune Origz Parka Hombre", "Color: Negro · Talle: L")
        _venta(db, "NADA-QUE-VER", "Producto raro")
    with TestClient(app) as c:
        _login(c)
        assert c.get("/articulos/cruce-sku").status_code == 200
        d = c.get("/articulos/api/cruce-sku").json()
        por = {f["sku"]: f for f in d["filas"]}
        assert por["M-777-BLACK-XL"]["estado"] == "automatico"
        assert por["M-777-BLACK-XL"]["articulo"]["sku"] == "1-CAMP-M-00777-BLCK-XL-PARK"   # Parka primero
        assert por["BUZO-VIEJO-1"]["estado"] == "sugerido"
        assert por["BUZO-VIEJO-1"]["articulo"]["sku"] == "1-BUZO-M-ORIGZ-BLCK-L-PARK"
        assert por["NADA-QUE-VER"]["estado"] == "sin_cruce"
        # confirmar a mano y quitar
        aid = por["BUZO-VIEJO-1"]["articulo"]["id"]
        assert c.post("/articulos/api/cruce-sku/confirmar", json={"sku": "buzo-viejo-1", "articulo_id": aid}).json()["ok"]
        por = {f["sku"]: f for f in c.get("/articulos/api/cruce-sku").json()["filas"]}
        assert por["BUZO-VIEJO-1"]["estado"] == "vinculado" and por["BUZO-VIEJO-1"]["por"] == "Vengil"
        assert c.post("/articulos/api/cruce-sku/quitar", json={"sku": "BUZO-VIEJO-1"}).json()["ok"]
        # confirmar todos los automáticos
        assert c.post("/articulos/api/cruce-sku/automaticos").json()["confirmados"] >= 1
        por = {f["sku"]: f for f in c.get("/articulos/api/cruce-sku").json()["filas"]}
        assert por["M-777-BLACK-XL"]["estado"] == "vinculado"
        # Excel de ida y vuelta
        x = c.get("/articulos/cruce-sku.xlsx")
        assert x.status_code == 200 and load_workbook(io.BytesIO(x.content)).active["A1"].value == "SKU ML / TN"
        wb = Workbook()
        wb.active.append(["SKU ML / TN", "SKU Odoo"])
        wb.active.append(["NADA-QUE-VER", "1-CAMP-M-00777-BLCK-XL-NOBR"])
        wb.active.append(["OTRO", "NO-EXISTE"])
        buf = io.BytesIO()
        wb.save(buf)
        r = c.post("/articulos/api/cruce-sku/importar", files={"archivo": ("c.xlsx", buf.getvalue())}).json()
        assert r["confirmados"] == 1 and len(r["errores"]) == 1
        por = {f["sku"]: f for f in c.get("/articulos/api/cruce-sku").json()["filas"]}
        assert por["NADA-QUE-VER"]["estado"] == "vinculado"
        assert por["NADA-QUE-VER"]["articulo"]["marca"] == "No Brand"
