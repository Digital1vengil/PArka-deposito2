import base64
import io
import re

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db import SessionLocal, hoy
from app.main import app
from app.models import Conteo, ConteoItem, DespachoEscaneo, Devolucion, DevolucionItem, Usuario
from app.services import deposito_drive, re_drive


def _login(c, nombre="Vengil", pin="1234"):
    uid = re.search(r"elegir\((\d+), '" + nombre + r"'\)", c.get("/login").text).group(1)
    assert c.post("/login", data={"usuario_id": uid, "pin": pin}).json()["ok"]


def _cargar_dia():
    with SessionLocal() as db:
        u = db.query(Usuario).first()
        for tipo, tr in [("flex", "44000000001"), ("flex", "4002"), ("colecta", "5001"), ("tiendanube", "TN-9")]:
            db.add(DespachoEscaneo(fecha=hoy(), tipo=tipo, tracking=tr, venta="V" + tr, comprador="Ana", usuario_id=u.id))
        d = Devolucion(fecha=hoy(), usuario_id=u.id, total=3, plataforma="ML")
        d.items = [DevolucionItem(codigo_barras="200", articulo="ABBA", color="BLACK", talle="M", estado="OPTIMO",
                                  plataforma="ML", venta="2000012345678901", llego="MISMO", cantidad=2),
                   DevolucionItem(codigo_barras="201", articulo="ABBA", color="BLACK", talle="L", estado="FALLA",
                                  plataforma="ML", cantidad=1)]
        c = Conteo(fecha=hoy(), usuario_id=u.id, nota="Estante A", total=5)
        c.items = [ConteoItem(codigo_barras="300", articulo="W13", color="NEGRO", talle="S", cantidad=5)]
        db.add_all([d, c])
        db.commit()


def _hojas(contenido: bytes) -> dict:
    wb = load_workbook(io.BytesIO(contenido))
    def fila(r):
        r = list(r)
        while r and r[-1] is None:
            r.pop()
        return r
    return {ws.title: [fila(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def test_excels_y_subida(monkeypatch):
    subidas = []

    def fake_guardar(db, fecha, archivos):
        subidas.append((fecha, archivos))
        return {"fecha": fecha, "carpeta": "https://drive/DEPOSITO",
                "archivos": [{"seccion": s, "nombre": f"{fecha} - {s.capitalize()}.xlsx", "url": f"https://drive/{s}"}
                             for s in archivos]}
    monkeypatch.setattr(re_drive, "configurado", lambda db=None: True)
    monkeypatch.setattr(re_drive, "guardar_deposito", fake_guardar)
    deposito_drive._subido.clear()

    with SessionLocal() as db:  # día vacío: el automático no sube nada
        assert deposito_drive.sincronizar(db, hoy())["archivos"] == []
    assert subidas == []

    _cargar_dia()
    with SessionLocal() as db:
        r = deposito_drive.sincronizar(db, hoy())
        assert {a["seccion"] for a in r["archivos"]} == {"despachos", "devoluciones", "conteos", "cierre"}
        # sin cambios: no se vuelve a subir
        assert deposito_drive.sincronizar(db, hoy())["archivos"] == []
    fecha, archivos = subidas[0]
    assert fecha == hoy().isoformat()

    desp = _hojas(archivos["despachos"])
    assert {"Resumen", "Flex Mercado Libre", "Flex Tienda Nube", "Colecta", "Tienda Nube"} <= set(desp)
    # en Flex, lo de Mercado Libre (envío de 11 dígitos) y lo de Tienda Nube (nº de orden) van en hojas separadas
    assert len(desp["Flex Mercado Libre"]) == 2 and desp["Flex Mercado Libre"][1][1] == "44000000001"
    assert len(desp["Flex Tienda Nube"]) == 2 and desp["Flex Tienda Nube"][1][1] == "4002"
    assert len(desp["Colecta"]) == 2 and len(desp["Tienda Nube"]) == 2
    assert ["Flex Mercado Libre", 1] in desp["Resumen"] and ["Flex Tienda Nube", 1] in desp["Resumen"]
    assert ["TOTAL", 4] in desp["Resumen"]

    dev = _hojas(archivos["devoluciones"])
    assert ["ML", 2, 1, 3] in dev["Resumen"] and len(dev["Detalle"]) == 3
    # después de "Plataforma" (D) va el nº de venta (ML) / orden (TN) en la columna E
    assert dev["Detalle"][0][3:5] == ["Plataforma", "Nº venta / orden"] and dev["Detalle"][1][3:5] == ["ML", "2000012345678901"]
    assert dev["Detalle"][0][-3] == "Llegó" and dev["Detalle"][1][-3] == "Mismo artículo"
    assert dev["Detalle"][0][-2:] == ["Ubicación", "Devuelto al rack"]

    con = _hojas(archivos["conteos"])
    assert con["Por artículo"][1][-1] == 5 and con["Detalle"][1][2] == "Estante A"

    cie = _hojas(archivos["cierre"])
    assert {"Resumen", "Despachados", "Devoluciones", "Conteos", "Tareas", "Actividad"} <= set(cie)
    assert ["Envíos despachados", 4] in cie["Resumen"] and len(cie["Despachados"]) == 5

    # un cambio en devoluciones vuelve a subir esa sección y el cierre, nada más
    with SessionLocal() as db:
        d = Devolucion(fecha=hoy(), total=1, plataforma="TN")
        d.items = [DevolucionItem(codigo_barras="202", articulo="ABBA", estado="OPTIMO", plataforma="TN", cantidad=1)]
        db.add(d)
        db.commit()
        r = deposito_drive.sincronizar(db, hoy())
    assert set(subidas[-1][1]) == {"devoluciones", "cierre"}

    # botón de Cierre diario: fuerza los 4
    with TestClient(app) as c:
        _login(c)
        assert "DEPOSITO PARKA" in c.get("/admin/cierre").text
        r = c.post("/admin/api/deposito-drive/guardar", json={}).json()
        assert r["ok"] and len(r["archivos"]) == 4
    assert set(subidas[-1][1]) == {"despachos", "devoluciones", "conteos", "cierre"}


def test_payload_guardar_deposito(monkeypatch):
    enviado = {}
    monkeypatch.setattr(re_drive, "_post", lambda db, action, timeout=120, **body: enviado.update(action=action, **body) or {})
    re_drive.guardar_deposito(None, "2026-10-16", {"conteos": b"abc"})
    assert enviado["action"] == "guardarDeposito" and enviado["fecha"] == "2026-10-16"
    assert enviado["archivos"] == [{"seccion": "conteos", "base64": base64.b64encode(b"abc").decode()}]


def test_cierre_por_mail(monkeypatch):
    mails = []
    monkeypatch.setattr(re_drive, "configurado", lambda db=None: True)
    monkeypatch.setattr(re_drive, "guardar_deposito", lambda db, fecha, archivos: {
        "archivos": [{"seccion": s, "nombre": s, "url": f"https://drive/{s}"} for s in archivos]})
    monkeypatch.setattr(re_drive, "enviar_mail", lambda db, para, asunto, html, adjuntos=None, copia="":
                        mails.append(dict(para=para, asunto=asunto, html=html, adjuntos=adjuntos)) or {})
    with SessionLocal() as db:
        assert not deposito_drive.cierre_mail_ya_enviado(db, hoy())
        r = deposito_drive.enviar_cierre_mail(db, hoy(), automatico=True)
        assert r.ok and r.destinatario == "gerencia@magontex.com.ar"
        assert deposito_drive.cierre_mail_ya_enviado(db, hoy())
    m = mails[0]
    assert m["para"] == "gerencia@magontex.com.ar" and "Cierre PARKA Depósito" in m["asunto"]
    assert "https://drive/cierre" in m["html"] and "Despachos" in m["html"]
    nombre, contenido = next(iter(m["adjuntos"].items()))
    assert nombre.endswith(" - Cierre.xlsx") and "Resumen" in _hojas(contenido)

    # si el script falla queda registrado el error
    def falla(*a, **k):
        raise re_drive.REDriveError("Acción desconocida: enviarMail")
    monkeypatch.setattr(re_drive, "enviar_mail", falla)
    with TestClient(app) as c:
        _login(c)
        r = c.post("/admin/api/deposito-drive/mail", json={})
        assert r.status_code == 502 and "enviarMail" in r.json()["detail"]
        assert "gerencia@magontex.com.ar" in c.get("/admin/cierre").text


def test_botones_manuales_usan_el_cierre_nuevo(monkeypatch):
    from app.services import drive
    llamadas = []
    monkeypatch.setattr(re_drive, "configurado", lambda db=None: True)
    monkeypatch.setattr(re_drive, "guardar_deposito", lambda db, fecha, archivos: {"archivos": []})
    monkeypatch.setattr(re_drive, "enviar_mail", lambda db, para, asunto, html, adjuntos=None, copia="":
                        llamadas.append(("mail_nuevo", para)) or {})
    monkeypatch.setattr(drive, "informar_dia", lambda db, u=None: {"despachados": [], "pendientes": []})
    monkeypatch.setattr(drive, "gas_post", lambda action, *a, **k: llamadas.append(("gas", action)) or {"ok": True})
    with TestClient(app) as c:
        _login(c)
        r = c.post("/despacho/api/cierre-mail").json()
        assert r["ok"] and r["mail"] == "gerencia@magontex.com.ar" and r["despachados"] >= 4
        assert c.post("/admin/api/google/cerrar").json()["ok"]
    assert llamadas == [("mail_nuevo", "gerencia@magontex.com.ar")] * 2   # nunca el cerrarDia viejo


def test_caracteres_invisibles_no_rompen_el_excel(monkeypatch):
    from app.services.deposito_drive import excel_despachos, excel_cierre_dia
    with SessionLocal() as db:
        u = db.query(Usuario).first()
        db.add(DespachoEscaneo(fecha=hoy(), tipo="flex", tracking="44000009001", venta="V9001",
                               comprador="Juan\x0b Pérez\x01", usuario_id=u.id))
        db.commit()
        hojas = _hojas(excel_despachos(db, hoy()))
        assert any(r and r[3] == "Juan Pérez" for r in hojas["Flex Mercado Libre"])
        assert "Resumen" in _hojas(excel_cierre_dia(db, hoy()))
