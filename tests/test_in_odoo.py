import io
import json
from datetime import datetime

from openpyxl import load_workbook

from app.models import Devolucion, DevolucionItem
from app.services import in_odoo


def _item(**k):
    base = dict(articulo="ORIGIN 00002", color="BLACK", cantidad=1, estado="OPTIMO", guardados="",
                devuelto_at=datetime(2026, 10, 21, 15, 42))
    base.update(k)
    return DevolucionItem(**base)


def _dev():
    d = Devolucion(id=15, plataforma="ML+TN")
    d.items = [
        _item(sku="1-CAMP-M-WORKJ-BLCK-XXL-PARK", codigo_barras="2000000011318", talle="XXL", plataforma="TN",
              guardados=json.dumps({"M1-R4": 1})),
        _item(sku="1-CAMP-M-WORKJ-BLCK-M-PARK", codigo_barras="2000000011226", talle="M", plataforma="ML",
              guardados=json.dumps({"M3-R2": 1})),
        _item(sku="1-CAMP-M-WORKJ-BLCK-L-PARK", codigo_barras="0000000011196", talle="L", plataforma="ML", cantidad=2,
              guardados=json.dumps({"M3-R2": 2})),
        _item(sku="1-CAMP-M-WORKJ-BLCK-S-PARK", codigo_barras="2000000011100", talle="S", plataforma="ML",
              estado="FALLA", guardados=json.dumps({"FALLAS": 1})),                 # FALLA: no va
        _item(sku="1-CAMP-M-WORKJ-BLCK-XL-PARK", codigo_barras="2000000011200", talle="XL", plataforma="ML"),  # sin guardar
        _item(sku="", codigo_barras="", talle="L", plataforma="ML", guardados=json.dumps({"M3-R2": 1})),  # sin SKU
    ]
    return d


def test_excel_in_odoo_formato():
    d = _dev()
    contenido, n, avisos = in_odoo.excel(d)
    assert n == 4 and len(avisos) == 1 and "no coincide" in avisos[0]
    assert in_odoo.nombre_archivo(d) == "IN_DEV_2026-10-21_PD-15.xlsx"
    wb = load_workbook(io.BytesIO(contenido))
    assert wb.sheetnames == ["IN DEVOLUCIONES"]
    ws = wb.active
    filas = [[c.value for c in r] for r in ws.iter_rows()]
    assert filas[0] == in_odoo.ENCABEZADOS and len(filas[0]) == 17 and filas[0][3] == "origin"
    assert len(filas) == 5
    # bloque ML primero (3 escaneos: M + L x2), después TN; partner unificado y origin = nombre del archivo
    assert filas[1][:9] == ["Magontex: Recepciones", "Devoluciones Ecommerce", "2026-10-21", "IN_DEV_2026-10-21_PD-15",
                            "Partners/Customers", "WH/Existencias", "1-CAMP-M-WORKJ-BLCK-M-PARK",
                            "1-CAMP-M-WORKJ-BLCK-M-PARK", 1]
    assert filas[1][9:] == ["PD-15", "ML", "2000000011226", "ORIGIN 00002", "BLACK", "M", "M3-R2", "OPTIMO"]
    assert filas[2][:6] == [None] * 6 and filas[3][:6] == [None] * 6
    assert filas[2][6] == filas[3][6] == "1-CAMP-M-WORKJ-BLCK-L-PARK" and filas[2][8] == 1
    assert filas[3][11] == "0000000011196" and ws["L4"].data_type == "s"   # código como texto, con sus ceros
    assert filas[4][:6] == ["Magontex: Recepciones", "Devoluciones Ecommerce", "2026-10-21", "IN_DEV_2026-10-21_PD-15",
                            "Partners/Customers", "WH/Existencias"]
    assert filas[4][9:] == ["PD-15", "TN", "2000000011318", "ORIGIN 00002", "BLACK", "XXL", "M1-R4", "OPTIMO"]
    assert ws["C2"].data_type == "s" and ws["C2"].number_format == "@"
    assert all(r[16] == "OPTIMO" for r in filas[1:])
    assert not any("MERCADO" in str(c).upper() or "TIENDA NUBE" in str(c).upper() for r in filas for c in r if c)


def test_sin_optimos_no_genera():
    d = Devolucion(id=16, plataforma="ML")
    d.items = [_item(sku="X-1", codigo_barras="1", plataforma="ML", estado="FALLA", guardados=json.dumps({"F": 1}))]
    contenido, n, avisos = in_odoo.excel(d)
    assert contenido is None and n == 0


def test_nombre_dev_correlativo():
    """Los archivos se llaman DEV-01, DEV-02… (correlativo único, no se reinicia)."""
    d = _dev()
    d.numero_dev = 7
    assert in_odoo.nombre_archivo(d) == "DEV-07.xlsx"
    d.numero_dev = 123
    assert in_odoo.codigo_dev(d) == "DEV-123"
    d.numero_dev = 1
    contenido, _n, _a = in_odoo.excel(d)
    ws = load_workbook(io.BytesIO(contenido)).active
    assert ws["D2"].value == "DEV-01" and ws["B2"].value == "Devoluciones Ecommerce"
