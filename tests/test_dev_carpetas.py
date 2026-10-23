"""Estructura de carpetas 1- RECEPCIONES … 4 - historial de in cargados a odoo (23/10/2026)."""
import io
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.models import Devolucion, DevolucionItem, Usuario
from app.services import dev_archivos, in_odoo


def _hoja(contenido):
    ws = load_workbook(io.BytesIO(contenido)).active
    return ws, [[c.value for c in r] for r in ws.iter_rows()]


def _item(**k):
    base = dict(articulo="ORIGIN 00002", color="BLACK", cantidad=1, estado="OPTIMO", guardados="", devuelto_cant=0,
                plataforma="ML", venta="2000009222", llego="MISMO")
    base.update(k)
    return DevolucionItem(**base)


def _dev(ruta="PENDIENTE"):
    d = Devolucion(id=40, fecha=date(2026, 10, 23), creado_at=datetime(2026, 10, 23, 9, 15), plataforma="ML+TN",
                   ruta_estado=ruta, numero_dev=3, in_estado="")
    d.usuario = Usuario(nombre="Diego")
    d.items = [
        _item(sku="1-CAMP-M-WORKJ-BLCK-M-PARK", codigo_barras="0000000011226", talle="M", cantidad=2,
              guardados=json.dumps({"B01-C07-N0": 2}), devuelto_cant=2, devuelto_at=datetime(2026, 10, 23, 10, 5),
              ubicacion="B01-C07-N0"),
        _item(sku="1-CAMP-M-WORKJ-BLCK-L-PARK", codigo_barras="2000000011196", talle="L", plataforma="TN", venta="1234",
              llego="DISTINTO", reemplaza="ORIGIN 00002 BLACK XL"),
        _item(sku="1-CAMP-M-WORKJ-BLCK-S-PARK", codigo_barras="2000000011100", talle="S", estado="FALLA"),
        _item(sku="", codigo_barras="", articulo="Buzo de otra marca", color="", talle="", estado="AJENO", llego="AJENO",
              reemplaza="ORIGIN 00002 BLACK XXL"),
    ]
    return d


def test_recepciones_registra_todo_con_estados():
    d = _dev()
    ws, f = _hoja(dev_archivos.excel_recepciones(d))
    assert f[0] == dev_archivos.ENC_1 and len(f) == 5   # Óptimo, Falla y "no es nuestro": todo lo que llegó
    col = {k: n for n, k in enumerate(dev_archivos.ENC_1)}
    m, l, s, x = f[1:]
    assert m[col["Fecha"]] == "2026-10-23" and m[col["Hora"]] == "09:15" and m[col["Usuario"]] == "Diego"
    assert m[col["Archivo"]] == "DEV-03" and m[col["Plataforma"]] == "Mercado Libre" and l[col["Plataforma"]] == "Tienda Nube"
    assert m[col["Código de barras"]] == "0000000011226" and ws["H2"].data_type == "s"   # sin perder ceros
    assert m[col["Validación"]] == "Mismo artículo" and l[col["Validación"]] == "Artículo distinto"
    assert l[col["Artículo de la venta"]] == "ORIGIN 00002 BLACK XL" and x[col["Validación"]] == "No es nuestro artículo"
    assert m[col["Estado"]] == "Devuelto" and l[col["Estado"]] == "Pendiente"
    assert s[col["Estado"]] == "Falla - no se ubica" and x[col["Estado"]] == "No es nuestro - no se ubica"
    # al cerrar la ruta lo que no se guardó pasa a "No devuelto"; "sin devolver" lo marca así
    d.ruta_estado = "DEVUELTO"
    assert dev_archivos.estado_item(d, d.items[1]) == "No devuelto"
    d.ruta_estado = "SIN_DEVOLVER"
    assert dev_archivos.estado_item(d, d.items[1]) == "Sin devolver"


def test_in_devol_ubicadas_es_la_fuente_del_in():
    d = _dev()
    ws, f = _hoja(dev_archivos.excel_ubicadas(d))
    assert f[0] == dev_archivos.ENC_2U
    assert f[1:] == [["2026-10-23", "10:05", "PD-40", "DEV-03", "Mercado Libre", "2000009222", "0000000011226",
                      "1-CAMP-M-WORKJ-BLCK-M-PARK", "ORIGIN 00002", "BLACK", "M", "B01-C07-N0", 2, "In-Devol-ubicadas"]]
    # el IN de "2- in odoo" tiene exactamente lo ubicado (con N° devolución y plataforma)
    d.ruta_estado = "DEVUELTO"
    _ws, fin = _hoja(in_odoo.excel(d)[0])
    assert len(fin) - 1 == sum(r[12] for r in f[1:]) and {(r[9], r[10]) for r in fin[1:]} == {("PD-40", "ML")}
    assert fin[1][3] == "DEV-03"
    d.in_estado = dev_archivos.IN_PENDIENTE
    assert {r[-1] for r in dev_archivos.filas_ubicadas(d)} == {"Pasado a 2- in odoo"}
    sin = _dev()
    for i in sin.items:
        i.guardados, i.devuelto_cant = "", 0
    assert dev_archivos.excel_ubicadas(sin) is None


def test_historial_diario_y_carpeta_4(monkeypatch):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.db import SessionLocal
    from app.services import re_drive, odoo
    subidos = []
    monkeypatch.setattr(re_drive, "guardar_deposito_archivo",
                        lambda db, sec, nombre, b: subidos.append((sec, nombre, b)) or {"url": "u"})
    monkeypatch.setattr("app.actividad.registrar", lambda *a, **k: None)
    with TestClient(app):
        with SessionLocal() as db:
            d = Devolucion(fecha=date(2026, 10, 19), creado_at=datetime(2026, 10, 19, 11, 30), plataforma="ML+TN",
                           ruta_estado="DEVUELTO")
            db.add(d)
            db.commit()
            for it in _dev().items:
                cols = ("sku", "codigo_barras", "articulo", "color", "talle", "estado", "plataforma", "venta", "llego",
                        "reemplaza", "guardados", "devuelto_cant", "devuelto_at", "ubicacion", "cantidad")
                db.add(DevolucionItem(devolucion_id=d.id, **{c: getattr(it, c) for c in cols}))
            # una prenda TN también ubicada → el IN tiene 2 bloques (ML y TN)
            db.commit()
            db.refresh(d)
            tn = next(i for i in d.items if i.plataforma == "TN")
            tn.guardados, tn.devuelto_cant, tn.devuelto_at = json.dumps({"A01-C01-N0": 1}), 1, datetime(2026, 10, 19, 12)
            db.commit()

            r = dev_archivos.generar_y_subir(db, d)
            assert r["ok"] and r["filas"] == 3 and r["nombre"] == in_odoo.nombre_archivo(d)
            claves = [s[0] for s in subidos]
            assert claves == ["rec_2_in_odoo", "rec_1_recepciones", "rec_2_ubicadas", "rec_3_historial", "rec_4_in_cargados"]
            assert not any(c in ("devoluciones", "credito_odoo") for c in claves)   # la carpeta vieja no se toca
            assert d.in_estado == "PENDIENTE_CARGA" and d.in_generado_at

            hist = next(s for s in subidos if s[0] == "rec_3_historial")
            assert hist[1] == "2026-10-19 - Historial Recepciones.xlsx"
            _ws, f3 = _hoja(hist[2])
            assert f3[0][:6] == ["Fecha", "Hora", "Artículo", "SKU", "Cantidad", "Origen"]
            assert ["2026-10-19", "11:30", "ORIGIN 00002 BLACK M", "1-CAMP-M-WORKJ-BLCK-M-PARK", 2, "Mercado Libre",
                    f"PD-{d.id}", "Devuelto"] in f3
            assert any(r[5] == "Tienda Nube" for r in f3[1:]) and len([r for r in f3[1:] if r[6] == f"PD-{d.id}"]) == 4

            origen = in_odoo.codigo_dev(d)
            fila4 = next(x for x in dev_archivos.filas_in_cargados(db) if x[0] == origen)
            assert fila4[1:6] == [f"PD-{d.id}", "Mercado Libre + Tienda Nube", 3, fila4[4], "Pendiente de carga"]

            # Odoo: se busca por ORIGEN. Con un solo IN (falta el de TN) sigue pendiente; con los dos, Cargado
            pickings = [{"name": "WH/IN/00950", "origin": origen, "state": "assigned"}]

            class Fake:
                def call(self, modelo, metodo, dominio, **kw):
                    assert modelo == "stock.picking" and dominio[0][0] == "origin" and origen in dominio[0][2]
                    return list(pickings)
            monkeypatch.setattr(odoo, "configurado", lambda db=None: True)
            monkeypatch.setattr(odoo, "cliente", lambda db=None, timeout=60: Fake())
            v = dev_archivos.verificar_in_odoo(db)
            assert origen not in v["cargados"] and d.in_estado == "PENDIENTE_CARGA" and d.in_odoo_ref == "WH/IN/00950"
            pickings.append({"name": "WH/IN/00951", "origin": origen, "state": "done"})
            subidos.clear()
            v = dev_archivos.verificar_y_subir(db)
            assert origen in v["cargados"] and d.in_estado == "CARGADO" and d.in_cargado_at
            assert [s[0] for s in subidos] == ["rec_4_in_cargados"]
            fila4 = next(x for x in dev_archivos.filas_in_cargados(db) if x[0] == origen)
            assert fila4[5] == "Cargado" and fila4[6] == "WH/IN/00950, WH/IN/00951"
            # volver a generar no lo devuelve a "pendiente de carga"
            dev_archivos.generar_y_subir(db, d)
            assert d.in_estado == "CARGADO"


def test_sin_devolver_no_genera_in(monkeypatch):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.db import SessionLocal
    from app.services import re_drive
    subidos = []
    monkeypatch.setattr(re_drive, "guardar_deposito_archivo", lambda db, sec, nombre, b: subidos.append(sec) or {"url": "u"})
    monkeypatch.setattr("app.actividad.registrar", lambda *a, **k: None)
    with TestClient(app):
        with SessionLocal() as db:
            d = Devolucion(fecha=date(2026, 10, 23), plataforma="ML", ruta_estado="SIN_DEVOLVER")
            db.add(d)
            db.commit()
            db.add(DevolucionItem(devolucion_id=d.id, sku="X-1", articulo="ABBA", estado="OPTIMO", plataforma="ML",
                                  venta="200001", cantidad=1))
            db.commit()
            db.refresh(d)
            r = dev_archivos.generar_y_subir(db, d)
            assert r["ok"] and not r["nombre"] and "no se genera IN" in r["motivo"] and not d.in_estado
            assert subidos == ["rec_1_recepciones", "rec_3_historial"]


def test_endpoints_carpeta_4(monkeypatch):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.services import re_drive
    from tests.test_devoluciones import _login
    monkeypatch.setattr(re_drive, "guardar_deposito_archivo", lambda db, sec, nombre, b: {"url": "u"})
    with TestClient(app) as c:
        _login(c)
        assert "revisados" in c.post("/devoluciones/api/in-cargados/verificar").json()
        x = c.get("/devoluciones/in-cargados.xlsx")
        assert x.status_code == 200 and "Historial IN cargados a Odoo.xlsx" in x.headers["content-disposition"]
        _ws, f = _hoja(x.content)
        assert f[0] == dev_archivos.ENC_4 and f[0][0] == "Origen"


def test_apps_script_carpetas_literales():
    gs = (Path(__file__).resolve().parents[1] / "google" / "deposito-drive.gs").read_text(encoding="utf-8")
    for clave, nombre in dev_archivos.CARPETAS.items():
        assert f"{clave}:" in gs and f"'{nombre}'" in gs
    assert "['In-Devol-ubicadas', 'rec_2_in_odoo']" in gs   # subcarpeta de "2- in odoo"
    secciones = gs[gs.index("var SECCIONES"):gs.index("};", gs.index("var SECCIONES"))]
    assert "devoluciones" not in secciones and "VERSION = '2026.10.23-1'" in gs
