"""Archivos de devoluciones ecommerce (ML + TN) en Drive — estructura de carpetas 1 a 4 (desde 23/10/2026).

Reemplaza a 01_Cantidades_Netas … 04_Carga_Odoo y a la carpeta "Devoluciones" (Excel del día): la app ya NO genera
ni guarda nada en la carpeta de devoluciones. Todo vive en DEPOSITO PARKA/Plantilla credito odoo/:

  1- RECEPCIONES/                    DEV-xx.xlsx  todo lo que llegó en la devolución (Óptimo, Falla y "no es nuestro")
  2- in odoo/                        DEV-xx.xlsx  plantilla del IN para "LA CARGA ODOO" (solo Óptimo ubicado)
  2- in odoo/In-Devol-ubicadas/      DEV-xx.xlsx  prendas ya ubicadas en su rack: de acá sale el IN de "2- in odoo"
  3- HISTORIAL RECEPCIONES/          AAAA-MM-DD - Historial Recepciones.xlsx   log diario de todo lo que llega
  4 - historial de in cargados a odoo/  Historial IN cargados a Odoo.xlsx       un renglón por ORIGEN (DEV-xx)

Flujo y estados
- Disparador: "Confirmar y armar la ruta". Ahí la devolución toma su N° DEV-xx y se generan 1 y 3 con todo en
  "Pendiente" (las Falla y lo que no es nuestro quedan como "no se ubica": NO entran a la ruta).
- Cada prenda Óptimo escaneada en su rack pasa a "Devuelto" en 1 y 3, y entra a In-Devol-ubicadas con estado
  "In-Devol-ubicadas". Se actualiza solo, en segundo plano, en cada escaneo.
- Al cerrar (última prenda o "Marcar como devuelto"): con lo de In-Devol-ubicadas se arma el IN en "2- in odoo"
  (origin = DEV-xx), esas filas pasan a "Pasado a 2- in odoo" y la devolución entra a 4 como "Pendiente de carga".
- 4 se verifica contra Odoo por el campo ORIGEN (stock.picking.origin = DEV-xx): aunque el IN tenga su propio
  número en Odoo, la referencia cruzada es siempre el origen. Si Odoo tiene la recepción → "Cargado".
"""
from __future__ import annotations

import io
import logging
import threading
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select

from ..models import Devolucion
from . import in_odoo

log = logging.getLogger("parka.dev_archivos")

PLATAFORMA_NOMBRE = {"ML": "Mercado Libre", "TN": "Tienda Nube"}
LLEGO = {"MISMO": "Mismo artículo", "DISTINTO": "Artículo distinto", "AJENO": "No es nuestro artículo"}

# clave que entiende el Apps Script (CARPETAS_DEV en deposito-drive.gs 2026.10.23-1) → nombre literal de la carpeta
CARPETAS = {
    "rec_1_recepciones": "1- RECEPCIONES",
    "rec_2_in_odoo": "2- in odoo",
    "rec_2_ubicadas": "In-Devol-ubicadas",              # subcarpeta de "2- in odoo"
    "rec_3_historial": "3- HISTORIAL RECEPCIONES",
    "rec_4_in_cargados": "4 - historial de in cargados a odoo",
}
ARCHIVO_4 = "Historial IN cargados a Odoo.xlsx"

# Estados (texto literal que se ve en los Excel)
PENDIENTE = "Pendiente"
DEVUELTO = "Devuelto"
UBICADA = "In-Devol-ubicadas"
PASADO_IN = "Pasado a 2- in odoo"
FALLA_NO_UBICA = "Falla - no se ubica"
AJENO_NO_UBICA = "No es nuestro - no se ubica"
SIN_DEVOLVER = "Sin devolver"
NO_DEVUELTO = "No devuelto"
IN_PENDIENTE = "PENDIENTE_CARGA"
IN_CARGADO = "CARGADO"
IN_TEXTO = {IN_PENDIENTE: "Pendiente de carga", IN_CARGADO: "Cargado"}

ENC_1 = ["Fecha", "Hora", "Usuario", "N° devolución", "Archivo", "Plataforma", "N° venta / orden", "Código de barras",
         "SKU", "Artículo", "Color", "Talle", "Cantidad", "Estado de la prenda", "Validación", "Artículo de la venta",
         "Estado", "Devueltas", "Ubicación", "Hora devuelto"]
ANCHO_1 = [11, 7, 14, 12, 9, 14, 20, 17, 32, 26, 14, 7, 9, 12, 22, 34, 26, 10, 18, 12]
ENC_2U = ["Fecha", "Hora", "N° devolución", "Origen", "Plataforma", "N° venta / orden", "Código de barras", "SKU",
          "Artículo", "Color", "Talle", "Rack", "Cantidad", "Estado"]
ANCHO_2U = [11, 7, 12, 9, 14, 20, 17, 32, 26, 14, 7, 16, 9, 22]
ENC_3 = ["Fecha", "Hora", "Artículo", "SKU", "Cantidad", "Origen", "N° devolución", "Estado"]
ANCHO_3 = [11, 7, 40, 32, 9, 14, 12, 24]
ENC_4 = ["Origen", "N° devolución", "Plataforma", "Prendas en el IN", "Generado", "Estado", "IN en Odoo",
         "Cargado (verificado)"]
ANCHO_4 = [10, 12, 26, 14, 17, 18, 30, 19]


# ─────────────────────────────── utilidades ───────────────────────────────
def _plat(i, d: Devolucion) -> str:
    return (i.plataforma or d.plataforma or "").strip().upper()


def _plat_nombre(p: str) -> str:
    return PLATAFORMA_NOMBRE.get(p, p)


def _desc(i) -> str:
    return " ".join(x for x in (i.articulo, i.color, i.talle) if x)


def _hhmm(dt) -> str:
    return dt.strftime("%H:%M") if dt else ""


def _libro(hoja: str, enc: list[str], anchos: list[int], filas: list[list], texto: tuple[str, ...] = ()) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = hoja[:31]
    ws.append(enc)
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in filas:
        ws.append(f)
        for col in texto:   # fechas, n° de venta y códigos como texto (no pierden ceros ni se vuelven número/fecha)
            c = ws[f"{col}{ws.max_row}"]
            c.number_format = "@"
            if c.value is not None:
                c.value, c.data_type = str(c.value), "s"
    for n, w in enumerate(anchos):
        ws.column_dimensions[chr(65 + n)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────────── estados ───────────────────────────────
def estado_item(d: Devolucion, i) -> str:
    """Estado de una prenda en los archivos. Cambia solo a medida que se escanea en los racks."""
    if i.estado == "AJENO":
        return AJENO_NO_UBICA
    if i.estado == "FALLA":
        return FALLA_NO_UBICA
    n, total = int(i.devuelto_cant or 0), int(i.cantidad or 0)
    if total and n >= total:
        return DEVUELTO
    if d.ruta_estado == "SIN_DEVOLVER":
        return SIN_DEVOLVER if not n else f"Devuelto parcial ({n} de {total}) · resto sin devolver"
    if d.ruta_estado == "DEVUELTO":
        return NO_DEVUELTO if not n else f"Devuelto parcial ({n} de {total}) · resto no devuelto"
    return PENDIENTE if not n else f"Devuelto parcial ({n} de {total})"


def validacion(i) -> str:
    """¿Llegó el mismo artículo que se vendió o uno distinto?"""
    if i.estado == "AJENO" or i.llego == "AJENO":
        return LLEGO["AJENO"]
    if i.llego == "DISTINTO":
        return LLEGO["DISTINTO"]
    return LLEGO["MISMO"]   # "MISMO" o sin marcar (carga manual del mismo artículo)


# ─────────────────────────────── 1- RECEPCIONES ───────────────────────────────
def filas_recepciones(d: Devolucion) -> list[list]:
    pd, dev = f"PD-{d.id}", in_odoo.codigo_dev(d)
    fecha = (d.creado_at.date() if d.creado_at else d.fecha) or date.today()
    usuario = d.usuario.nombre if d.usuario else ""
    out = []
    for i in d.items:
        out.append([fecha.isoformat(), _hhmm(d.creado_at), usuario, pd, dev, _plat_nombre(_plat(i, d)), i.venta or "",
                    i.codigo_barras or "", i.sku or "", i.articulo or "", i.color or "", i.talle or "",
                    int(i.cantidad or 0), "NO ES NUESTRO" if i.estado == "AJENO" else i.estado, validacion(i),
                    i.reemplaza or ("" if i.llego != "DISTINTO" else "—"), estado_item(d, i),
                    int(i.devuelto_cant or 0) if i.estado == "OPTIMO" else "", i.ubicacion or "",
                    _hhmm(i.devuelto_at) if i.estado == "OPTIMO" else ""])
    return out


def excel_recepciones(d: Devolucion) -> bytes:
    return _libro("RECEPCIONES", ENC_1, ANCHO_1, filas_recepciones(d), texto=("A", "G", "H"))


# ─────────────────────────────── 2- in odoo / In-Devol-ubicadas ───────────────────────────────
def filas_ubicadas(d: Devolucion) -> list[list]:
    """Una fila por prenda y rack donde quedó ubicada (Óptimo). Es la fuente del IN de "2- in odoo"."""
    pd, dev = f"PD-{d.id}", in_odoo.codigo_dev(d)
    from .ruta_devolucion import guardados
    estado = PASADO_IN if d.in_estado else UBICADA
    out = []
    for i in d.items:
        if i.estado != "OPTIMO":
            continue
        for rack, n in guardados(i).items():
            f = i.devuelto_at
            out.append([f.date().isoformat() if f else "", _hhmm(f), pd, dev, _plat_nombre(_plat(i, d)), i.venta or "",
                        i.codigo_barras or "", i.sku or "", i.articulo or "", i.color or "", i.talle or "", rack, int(n),
                        estado])
    return out


def excel_ubicadas(d: Devolucion) -> bytes | None:
    filas = filas_ubicadas(d)
    return _libro("IN-DEVOL-UBICADAS", ENC_2U, ANCHO_2U, filas, texto=("A", "F", "G")) if filas else None


# ─────────────────────────────── 3- HISTORIAL RECEPCIONES (diario) ───────────────────────────────
def nombre_historial(fecha: date) -> str:
    return f"{fecha.isoformat()} - Historial Recepciones.xlsx"


def devoluciones_del_dia(db, fecha: date) -> list[Devolucion]:
    return list(db.scalars(select(Devolucion).where(Devolucion.fecha == fecha).order_by(Devolucion.creado_at)).all())


def filas_historial(devs: list[Devolucion]) -> list[list]:
    out = []
    for d in devs:
        for i in d.items:
            out.append([(d.creado_at.date() if d.creado_at else d.fecha).isoformat(), _hhmm(d.creado_at),
                        _desc(i) or "(sin descripción)", i.sku or "", int(i.cantidad or 0), _plat_nombre(_plat(i, d)),
                        f"PD-{d.id}", estado_item(d, i)])
    return out


def excel_historial(db, fecha: date) -> bytes:
    return _libro("HISTORIAL RECEPCIONES", ENC_3, ANCHO_3, filas_historial(devoluciones_del_dia(db, fecha)),
                  texto=("A",))


# ─────────────────────────────── 4 - historial de in cargados a odoo ───────────────────────────────
def filas_in_cargados(db) -> list[list]:
    devs = db.scalars(select(Devolucion).where(Devolucion.in_estado != "").order_by(Devolucion.numero_dev)).all()
    out = []
    for d in devs:
        escaneos, _ = in_odoo.filas(d)
        plats = sorted({e["plataforma"] for e in escaneos}, key=lambda p: in_odoo.PLATAFORMAS.index(p))
        out.append([in_odoo.codigo_dev(d), f"PD-{d.id}", " + ".join(_plat_nombre(p) for p in plats), len(escaneos),
                    d.in_generado_at.strftime("%Y-%m-%d %H:%M") if d.in_generado_at else "",
                    IN_TEXTO.get(d.in_estado, d.in_estado), d.in_odoo_ref or "",
                    d.in_cargado_at.strftime("%Y-%m-%d %H:%M") if d.in_cargado_at else ""])
    return out


def excel_in_cargados(db) -> bytes:
    return _libro("HISTORIAL IN ODOO", ENC_4, ANCHO_4, filas_in_cargados(db), texto=("E", "H"))


def verificar_in_odoo(db) -> dict:
    """Busca en Odoo cada IN "Pendiente de carga" por su ORIGEN (DEV-xx). Hay un IN por plataforma (bloque ML y
    bloque TN), así que queda "Cargado" cuando Odoo tiene tantas recepciones con ese origen como bloques tiene.
    Devuelve {revisados, cargados: [DEV-xx], error?}."""
    from ..db import ahora
    from . import odoo
    pend = db.scalars(select(Devolucion).where(Devolucion.in_estado == IN_PENDIENTE)).all()
    r = {"revisados": len(pend), "cargados": []}
    if not pend or not odoo.configurado(db):
        return r
    por_origen = {in_odoo.codigo_dev(d): d for d in pend if in_odoo.codigo_dev(d)}
    try:
        recs = odoo.cliente(db, timeout=30).call(
            "stock.picking", "search_read", [["origin", "in", list(por_origen)], ["state", "!=", "cancel"]],
            fields=["name", "origin", "state"])
    except Exception as e:  # noqa: BLE001 — se reintenta en la próxima vuelta
        log.warning("No se pudo verificar los IN en Odoo: %s", e)
        r["error"] = str(e)[:300]
        return r
    encontrados: dict[str, list[str]] = {}
    for x in recs:
        encontrados.setdefault((x.get("origin") or "").strip(), []).append(x.get("name") or "")
    for origen, d in por_origen.items():
        nombres = encontrados.get(origen, [])
        if not nombres:
            continue
        bloques = len({e["plataforma"] for e in in_odoo.filas(d)[0]}) or 1
        d.in_odoo_ref = ", ".join(sorted(nombres))[:160]
        if len(nombres) >= bloques:
            d.in_estado, d.in_cargado_at = IN_CARGADO, ahora()
            r["cargados"].append(origen)
    db.commit()
    return r


# ─────────────────────────────── subida a Drive ───────────────────────────────
def _subir(db, clave: str, nombre: str, contenido: bytes) -> dict:
    from . import re_drive
    return re_drive.guardar_deposito_archivo(db, clave, nombre, contenido)


def sincronizar_estado(db, d: Devolucion) -> dict:
    """Sube lo que cambia en cada paso: 1- RECEPCIONES, In-Devol-ubicadas y el 3- HISTORIAL del día.
    Nunca toca la carpeta de devoluciones."""
    in_odoo.asegurar_numero(db, d)
    nombre = in_odoo.nombre_archivo(d)
    out = {"ok": True, "archivos": [], "errores": []}
    tareas = [("rec_1_recepciones", nombre, excel_recepciones(d))]
    ub = excel_ubicadas(d)
    if ub is not None:
        tareas.append(("rec_2_ubicadas", nombre, ub))
    fecha = d.fecha or (d.creado_at.date() if d.creado_at else date.today())
    tareas.append(("rec_3_historial", nombre_historial(fecha), excel_historial(db, fecha)))
    for clave, nom, contenido in tareas:
        try:
            res = _subir(db, clave, nom, contenido)
            out["archivos"].append({"clave": clave, "carpeta": CARPETAS[clave], "nombre": nom, "url": res.get("url", "")})
        except Exception as e:  # noqa: BLE001 — se reintenta en el próximo cambio o con "☁ Archivos a Drive"
            log.warning("PD-%s: no se pudo subir %s/%s: %s", d.id, CARPETAS[clave], nom, e)
            out["errores"].append(f"{CARPETAS[clave]}/{nom}: {e}")
    out["ok"] = not out["errores"]
    return out


def subir_historial_in(db) -> dict:
    return _subir(db, "rec_4_in_cargados", ARCHIVO_4, excel_in_cargados(db))


def generar_y_subir(db, d: Devolucion, usuario=None) -> dict:
    """Cierre de la devolución (o "☁ Archivos a Drive"): IN en "2- in odoo" a partir de In-Devol-ubicadas, alta en
    "4 - historial…" como Pendiente de carga y actualización de 1, In-Devol-ubicadas y 3.
    Devuelve {ok, nombre (el IN), filas, avisos, url, archivos: [...] | error | motivo} — mismas claves que antes."""
    from ..actividad import registrar
    from ..db import ahora
    r = {"ok": True, "nombre": "", "filas": 0, "avisos": [], "url": "", "archivos": []}
    try:
        in_odoo.asegurar_numero(db, d)
        contenido, n, avisos = (None, 0, [])
        if d.ruta_estado == "DEVUELTO":
            contenido, n, avisos = in_odoo.excel(d)
        r["avisos"] = list(avisos)
    except Exception as e:  # noqa: BLE001
        log.exception("No se pudo armar el IN de PD-%s", d.id)
        return dict(r, ok=False, error=f"No se pudo armar el IN de la devolución: {e}"[:300])
    errores = []
    if contenido is not None:
        nombre = in_odoo.nombre_archivo(d)
        if not d.in_estado:   # entra a 4 la primera vez que se genera; si ya estaba cargado no se pisa
            d.in_estado, d.in_generado_at = IN_PENDIENTE, ahora()
            db.commit()
        try:
            res = _subir(db, "rec_2_in_odoo", nombre, contenido)
            r.update(nombre=nombre, filas=n, url=res.get("url", ""))
            r["archivos"].append({"clave": "rec_2_in_odoo", "carpeta": CARPETAS["rec_2_in_odoo"], "nombre": nombre,
                                  "filas": n, "url": r["url"]})
        except Exception as e:  # noqa: BLE001
            errores.append(f"{CARPETAS['rec_2_in_odoo']}/{nombre}: {e}")
            r.update(nombre=nombre, filas=n)
    elif d.ruta_estado == "DEVUELTO":
        r["motivo"] = "No hay prendas OPTIMO ubicadas en rack: no se genera IN (sí 1- RECEPCIONES y el historial)"
    else:
        r["motivo"] = "Devolución sin ubicar: no se genera IN (sí 1- RECEPCIONES y el historial)"
    est = sincronizar_estado(db, d)
    r["archivos"] += est["archivos"]
    errores += est["errores"]
    if d.in_estado:
        try:
            res = subir_historial_in(db)
            r["archivos"].append({"clave": "rec_4_in_cargados", "carpeta": CARPETAS["rec_4_in_cargados"],
                                  "nombre": ARCHIVO_4, "url": res.get("url", "")})
        except Exception as e:  # noqa: BLE001
            errores.append(f"{CARPETAS['rec_4_in_cargados']}/{ARCHIVO_4}: {e}")
    if errores:
        r.update(ok=False, error=("No se pudo subir a Drive: " + " · ".join(errores))[:300])
    if r["archivos"]:
        try:
            registrar(db, usuario, "drive_deposito", f"PD-{d.id} ({in_odoo.codigo_dev(d)}): {len(r['archivos'])} archivos"
                      + (f" · IN {r['filas']} prendas" if r["filas"] else ""), "devoluciones")
        except Exception:  # noqa: BLE001
            log.warning("No se pudo registrar la actividad de PD-%s", d.id)
    return r


def verificar_y_subir(db) -> dict:
    """Job periódico: verifica los IN contra Odoo por ORIGEN y, si alguno pasó a Cargado, re-sube el archivo 4."""
    r = verificar_in_odoo(db)
    if r["cargados"]:
        try:
            subir_historial_in(db)
        except Exception as e:  # noqa: BLE001
            r["error"] = f"No se pudo subir {ARCHIVO_4}: {e}"[:300]
    return r


# ─────────────── en segundo plano: cada cambio de estado se refleja solo en los archivos ───────────────
_lock = threading.Lock()
_cola = {"corriendo": False, "ids": set()}


def programar(devolucion_id: int) -> None:
    """Encola la actualización de 1 / In-Devol-ubicadas / 3 de una devolución (no frena el escaneo)."""
    with _lock:
        _cola["ids"].add(devolucion_id)
        if _cola["corriendo"]:
            return
        _cola["corriendo"] = True
    threading.Thread(target=_procesar, daemon=True).start()


def _procesar() -> None:
    from ..db import SessionLocal
    from . import re_drive
    while True:
        with _lock:
            if not _cola["ids"]:
                _cola["corriendo"] = False
                return
            did = _cola["ids"].pop()
        try:
            with SessionLocal() as db:
                d = db.get(Devolucion, did)
                if d and re_drive.configurado(db):
                    sincronizar_estado(db, d)
        except Exception:  # noqa: BLE001
            log.exception("No se pudieron actualizar los archivos de PD-%s", did)


def zip_(d: Devolucion, db=None) -> tuple[bytes | None, str]:
    """Los archivos de la devolución en un .zip con la misma estructura de carpetas (para bajarlos desde la app)."""
    import zipfile
    nombre = in_odoo.nombre_archivo(d)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{CARPETAS['rec_1_recepciones']}/{nombre}", excel_recepciones(d))
        ub = excel_ubicadas(d)
        if ub is not None:
            z.writestr(f"{CARPETAS['rec_2_in_odoo']}/{CARPETAS['rec_2_ubicadas']}/{nombre}", ub)
        if d.ruta_estado == "DEVUELTO":
            contenido, _n, _a = in_odoo.excel(d)
            if contenido is not None:
                z.writestr(f"{CARPETAS['rec_2_in_odoo']}/{nombre}", contenido)
        if db is not None and d.fecha:
            z.writestr(f"{CARPETAS['rec_3_historial']}/{nombre_historial(d.fecha)}", excel_historial(db, d.fecha))
    cod = in_odoo.codigo_dev(d)
    return buf.getvalue(), (f"{cod}.zip" if cod else f"DEV_PD-{d.id}.zip")
