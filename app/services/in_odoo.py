"""Excel para el IN de devoluciones en Odoo (paso 2: prendas guardadas en su rack).

Se arma al confirmar la devolución en los racks (última prenda guardada o "Marcar como devuelto") y se sube a
DEPOSITO PARKA/Plantilla credito odoo/IN_DEV_AAAA-MM-DD_PD-<id>.xlsx. Después se importa en Odoo
(Inventario > Recepciones > Importar registros) y se crea un IN por plataforma.

Reglas (plantilla "Plantilla_IN_Devoluciones_ML_TN_ODOO"):
- Una sola hoja, encabezados exactos (Odoo los reconoce por nombre). Valores fijos, sin filas vacías ni totales.
- Una fila por escaneo del paso 2 (1 unidad). Solo prendas OPTIMO; las FALLA no van.
- Cada escaneo tiene que cruzar con una fila OPTIMO del paso 1 (N° devolución + SKU). Si no cruza, no se exporta
  y se avisa.
- Un bloque por plataforma: primero ML, abajo TN. A–F solo en la primera fila de cada bloque.
- partner_id unificado: "Devoluciones Ecommerce" para ML y TN (contacto existente en Odoo, id 13979; nombre exacto).
- origin (D) = nombre del archivo sin .xlsx (DEV-01, DEV-02…) → en Odoo cada IN queda con el archivo en "Documento origen".
- Nombre de los archivos (desde 22/10): DEV-<n°> correlativo único (DEV-01, DEV-02… DEV-100), se asigna una sola vez
  por devolución con asegurar_numero(); antes era IN_DEV_AAAA-MM-DD_PD-<id>.
- A–I se importan en Odoo; J–Q son control interno (al importar: "No importar").
"""
from __future__ import annotations

import io
import logging
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font

from ..models import Devolucion

log = logging.getLogger("parka.in_odoo")

ENCABEZADOS = [
    "picking_type_id", "partner_id", "scheduled_date", "origin", "location_id", "location_dest_id",
    "move_ids_without_package/product_id", "move_ids_without_package/name", "move_ids_without_package/product_uom_qty",
    "N° DEVOLUCIÓN", "PLATAFORMA", "CÓDIGO DE BARRAS", "ARTÍCULO", "COLOR", "TALLE", "UBICACIÓN", "ESTADO",
]
ANCHOS = [22, 22, 13, 26, 18, 16, 32, 32, 11, 12, 10, 17, 16, 10, 7, 14, 10]
TIPO_OPERACION = "Magontex: Recepciones"
CONTACTO = "Devoluciones Ecommerce"  # mismo contacto para ML y TN (no "Devolución Mercado Libre / Tienda Nube")
ORIGEN = "Partners/Customers"
DESTINO = "WH/Existencias"
PLATAFORMAS = ("ML", "TN")   # orden de los bloques
CARPETA = "credito_odoo"     # DEPOSITO PARKA/Plantilla credito odoo (Apps Script 2026.10.21-2)


def _guardados(i) -> dict[str, int]:
    from .ruta_devolucion import guardados
    return guardados(i)


def fecha_paso2(d: Devolucion) -> date:
    """Fecha del paso 2: la última prenda confirmada en un rack (si no hay, hoy)."""
    from ..db import hoy
    ts = [i.devuelto_at for i in d.items if getattr(i, "devuelto_at", None)]
    return max(ts).date() if ts else hoy()


def codigo_dev(d: Devolucion) -> str:
    """DEV-01, DEV-02… (vacío si todavía no tiene número)."""
    n = int(getattr(d, "numero_dev", 0) or 0)
    return f"DEV-{n:02d}" if n > 0 else ""


def asegurar_numero(db, d: Devolucion) -> str:
    """Le da a la devolución su N° correlativo DEV-xx la primera vez que se arman sus archivos (después no cambia)."""
    if not codigo_dev(d) and db is not None:
        from sqlalchemy import func
        ultimo = db.query(func.max(Devolucion.numero_dev)).scalar() or 0
        d.numero_dev = int(ultimo) + 1
        db.commit()
    return codigo_dev(d)


def nombre_archivo(d: Devolucion, fecha: date | None = None) -> str:
    """DEV-01.xlsx (sin número todavía: el nombre viejo IN_DEV_AAAA-MM-DD_PD-<id>.xlsx)."""
    return f"{codigo_dev(d)}.xlsx" if codigo_dev(d) else f"IN_DEV_{(fecha or fecha_paso2(d)).isoformat()}_PD-{d.id}.xlsx"


def filas(d: Devolucion) -> tuple[list[dict], list[str]]:
    """(escaneos que van al IN, avisos). Un escaneo = 1 unidad guardada en un rack en el paso 2."""
    pd = f"PD-{d.id}"
    # paso 1: filas OPTIMO por (N° devolución, SKU)
    paso1 = {(pd, (i.sku or "").strip().upper()) for i in d.items if i.estado == "OPTIMO" and (i.sku or "").strip()}
    out, avisos = [], []
    for i in d.items:
        g = _guardados(i)
        if not g or i.estado != "OPTIMO":       # FALLA (va a scrap, otro circuito) o AJENO: no van al IN
            continue
        sku = (i.sku or "").strip()
        plat = (i.plataforma or d.plataforma or "").strip().upper()
        nombre = " ".join(x for x in (i.articulo, i.color, i.talle) if x)
        if (pd, sku.upper()) not in paso1:
            avisos.append(f"{nombre or 'Una prenda'}: el escaneo no coincide con una fila OPTIMO de {pd} "
                          f"(SKU {sku or 'vacío'}) — no se exporta")
            continue
        if plat not in PLATAFORMAS:
            avisos.append(f"{nombre} ({sku}): no tiene plataforma ML o TN — no se exporta")
            continue
        for rack, n in g.items():
            for _ in range(int(n)):
                out.append({"_item": i, "venta": (i.venta or "").strip(),
                            "fecha": i.devuelto_at.date() if i.devuelto_at else None,
                            "sku": sku, "plataforma": plat, "devolucion": pd, "codigo": str(i.codigo_barras or ""),
                            "articulo": i.articulo or "", "color": i.color or "", "talle": i.talle or "",
                            "ubicacion": rack})
    return out, avisos


def excel(d: Devolucion, fecha: date | None = None) -> tuple[bytes | None, int, list[str]]:
    """(contenido .xlsx o None si no hay nada para el IN, filas exportadas, avisos)."""
    escaneos, avisos = filas(d)
    if not escaneos:
        return None, 0, avisos
    f = (fecha or fecha_paso2(d)).isoformat()
    origin = nombre_archivo(d, fecha)[:-5]   # DEV-01
    wb = Workbook()
    ws = wb.active
    ws.title = "IN DEVOLUCIONES"
    ws.append(ENCABEZADOS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for plat in PLATAFORMAS:
        bloque = [e for e in escaneos if e["plataforma"] == plat]
        for n, e in enumerate(bloque):
            cab = [TIPO_OPERACION, CONTACTO, f, origin, ORIGEN, DESTINO] if n == 0 else [None] * 6
            ws.append(cab + [e["sku"], e["sku"], 1, e["devolucion"], plat, e["codigo"], e["articulo"], e["color"],
                             e["talle"], e["ubicacion"], "OPTIMO"])
            fila = ws.max_row
            # fecha y código de barras como texto (no se pierden dígitos ni se convierte en fecha)
            for col in ("C", "L"):
                ws[f"{col}{fila}"].number_format = "@"
                if ws[f"{col}{fila}"].value is not None:
                    ws[f"{col}{fila}"].data_type = "s"
    for idx, w in enumerate(ANCHOS):
        ws.column_dimensions[chr(65 + idx)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(escaneos), avisos


def generar_y_subir(db, d: Devolucion, usuario=None) -> dict:
    """Arma el Excel del IN y lo sube a DEPOSITO PARKA/Plantilla credito odoo.
    Nunca corta el cierre de la devolución: devuelve {ok, nombre, filas, avisos, url | error}."""
    from ..actividad import registrar
    from . import re_drive
    try:
        contenido, n, avisos = excel(d)
    except Exception as e:  # noqa: BLE001
        log.exception("No se pudo armar el IN de PD-%s", d.id)
        return {"ok": False, "nombre": "", "filas": 0, "avisos": [], "error": f"No se pudo armar el Excel del IN: {e}"[:300]}
    nombre = nombre_archivo(d)
    r = {"ok": True, "nombre": nombre, "filas": n, "avisos": avisos, "url": ""}
    if contenido is None:
        r.update(ok=True, nombre="", motivo="No hay prendas OPTIMO guardadas en rack: no se genera IN")
        return r
    try:
        res = re_drive.guardar_deposito_archivo(db, CARPETA, nombre, contenido)
        r["url"] = res.get("url", "")
        registrar(db, usuario, "drive_deposito", f"{nombre} · {n} prendas para el IN de Odoo", "devoluciones")
    except Exception as e:  # noqa: BLE001 — se puede bajar desde la app y reintentar
        log.warning("IN de PD-%s: no se pudo subir a Drive: %s", d.id, e)
        r.update(ok=False, error=f"No se pudo subir {nombre} a Drive: {e}"[:300])
    return r
