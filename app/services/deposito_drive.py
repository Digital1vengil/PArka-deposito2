"""DEPOSITO PARKA en Google Drive: un Excel por día en cada carpeta.

  DEPOSITO PARKA/Despachos/     <fecha> - Despachos.xlsx     Flex + Colecta + Tienda Nube (y cambios / SmartPost)
  DEPOSITO PARKA/Devoluciones/  <fecha> - Devoluciones.xlsx  lo que se devolvió
  DEPOSITO PARKA/Conteos/       <fecha> - Conteos.xlsx       lo que se contó
  DEPOSITO PARKA/Cierre/        <fecha> - Cierre.xlsx        todo lo realizado y despachado en el día

La app arma los Excel con lo que hay en la base y el Apps Script (google/deposito-drive.gs, el mismo de
Recolección Ecommerce) solo los guarda: cada vez que se vuelve a mandar un día, el archivo se reemplaza.
Se actualizan solos cada 15 minutos si algo cambió (Conexiones → deposito_drive_auto) y a mano desde Cierre diario.

El cierre también sale por mail (cierre_mail_para, por defecto gerencia@magontex.com.ar) en los días y la hora
de cierre_mail_dias / cierre_mail_hora (lunes a viernes, 17:00), con el Excel del Cierre adjunto.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
from collections import Counter, defaultdict
from datetime import date, datetime, time

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..actividad import ACCIONES, registrar
from ..db import ahora
from ..models import (Conteo, DespachoEscaneo, Devolucion, EnvioManual, Lote, Paquete, Usuario, TIPOS_LOTE)
from . import ajustes, re_drive
from .cierre import datos_cierre

log = logging.getLogger("parka.deposito_drive")

SECCIONES = ("despachos", "devoluciones", "conteos", "cierre")
TIPOS_DESPACHO = {"flex": "Flex", "colecta": "Colecta", "tiendanube": "Tienda Nube", "cambios": "Cambios"}

# huella de lo último subido por (fecha, sección): si no cambió nada, no se vuelve a subir
_subido: dict[tuple[str, str], str] = {}


class DepositoDriveError(Exception):
    pass


def configurado(db: Session | None = None) -> bool:
    return re_drive.configurado(db)


def automatico(db: Session | None = None) -> bool:
    return ajustes.leer("deposito_drive_auto", db).strip() != "0"


# ─────────────── Datos del día ───────────────

def _rango(fecha: date) -> tuple[datetime, datetime]:
    return datetime.combine(fecha, time.min), datetime.combine(fecha, time.max)


def _hhmm(dt: datetime | None) -> str:
    return dt.strftime("%H:%M") if dt else ""


def despachos_del_dia(db: Session, fecha: date) -> list[dict]:
    """Lo que salió en el día: despacho en vivo (scanner) + paquetes de lotes escaneados ese día.
    Si un envío figura en los dos, se cuenta una sola vez."""
    inicio, fin = _rango(fecha)
    nombres = {u.id: u.nombre for u in db.scalars(select(Usuario)).all()}
    filas: dict[tuple[str, str], dict] = {}
    for r in db.scalars(select(DespachoEscaneo).where(DespachoEscaneo.fecha == fecha)
                        .order_by(DespachoEscaneo.at)).all():
        filas[(r.tipo, r.tracking)] = dict(tipo=r.tipo, hora=_hhmm(r.at), tracking=r.tracking, venta=r.venta,
                                           comprador=r.comprador, producto="", sku="", cantidad=None,
                                           operario=nombres.get(r.usuario_id, ""), origen="Scanner en vivo", _ts=r.at)
    paquetes = db.scalars(select(Paquete).join(Lote).where(Paquete.estado == "despachado",
                                                           Paquete.escaneado_at >= inicio,
                                                           Paquete.escaneado_at <= fin)).all()
    for p in paquetes:
        tipo = p.lote.tipo if p.lote.tipo in TIPOS_DESPACHO else "flex"
        k = (tipo, p.tracking)
        previo = filas.get(k)
        fila = dict(tipo=tipo, hora=_hhmm(p.escaneado_at), tracking=p.tracking, venta=p.venta_id,
                    comprador=p.comprador, producto=p.producto, sku=p.sku, cantidad=p.cantidad,
                    operario=p.escaneado_por.nombre if p.escaneado_por else "", origen=f"Lote: {p.lote.nombre}",
                    _ts=p.escaneado_at)
        if previo:  # el lote trae más detalle (producto / SKU): se completa la fila del scanner
            for c in ("venta", "comprador", "operario"):
                fila[c] = previo[c] or fila[c]
            fila["hora"], fila["_ts"] = previo["hora"], previo["_ts"]
        filas[k] = fila
    return sorted(filas.values(), key=lambda f: (f["_ts"] or datetime.min, f["tracking"]))


def smartpost_del_dia(db: Session, fecha: date) -> list[EnvioManual]:
    return list(db.scalars(select(EnvioManual).where(EnvioManual.fecha == fecha).order_by(EnvioManual.hora)).all())


def devoluciones_del_dia(db: Session, fecha: date) -> list[Devolucion]:
    return list(db.scalars(select(Devolucion).where(Devolucion.fecha == fecha).order_by(Devolucion.creado_at)).all())


def conteos_del_dia(db: Session, fecha: date) -> list[Conteo]:
    return list(db.scalars(select(Conteo).where(Conteo.fecha == fecha).order_by(Conteo.creado_at)).all())


# ─────────────── Excel ───────────────

_ENC_FONT = Font(bold=True, color="FFFFFF")
_ENC_FILL = PatternFill("solid", fgColor="1F2937")


def _limpio(v):
    """Excel no acepta caracteres de control (vienen a veces en nombres de compradores o detalles): se sacan."""
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)[:32000]
    return v


def _fila(ws, valores) -> None:
    ws.append([_limpio(v) for v in valores])


def _hoja(wb: Workbook, titulo: str, encabezados: list[str], filas: list[list], anchos: list[int], primera=False):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo[:31]
    _fila(ws, encabezados)
    for c in ws[1]:
        c.font, c.fill, c.alignment = _ENC_FONT, _ENC_FILL, Alignment(vertical="center")
    for f in filas:
        _fila(ws, f)
    for i, w in enumerate(anchos):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _titulo(ws, texto: str, sub: str = ""):
    ws.append([texto])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([sub])
    ws["A2"].font = Font(color="6B7280")
    ws.append([])


def _resumen(wb: Workbook, titulo: str, fecha: date, bloques: list[tuple[list[str], list[list]]]):
    ws = wb.active
    ws.title = "Resumen"
    _titulo(ws, f"{titulo} — {fecha.strftime('%d/%m/%Y')}", f"PARKA Depósito · actualizado {ahora().strftime('%d/%m %H:%M')}")
    for enc, filas in bloques:
        _fila(ws, enc)
        for c in ws[ws.max_row]:
            c.font, c.fill = _ENC_FONT, _ENC_FILL
        for f in filas:
            _fila(ws, f)
        ws.append([])
    for col, w in zip("ABCDE", [30, 16, 16, 16, 16]):
        ws.column_dimensions[col].width = w
    return ws


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ENC_DESP = ["Hora", "Nº de envío", "Venta", "Comprador", "Producto", "SKU", "Cant.", "Operario", "Origen"]
ANCHO_DESP = [8, 20, 18, 26, 40, 18, 7, 14, 28]
ENC_SP = ["Hora", "Plataforma", "Nombre", "Artículo", "Nº guía", "Guardado en Drive"]
ANCHO_SP = [8, 10, 26, 36, 20, 16]
ENC_DEV = ["Hora", "Devolución", "Responsable", "Plataforma", "Nº venta / orden", "Código", "SKU", "Artículo", "Color", "Talle",
           "Estado", "Cantidad"]
ANCHO_DEV = [8, 12, 16, 10, 20, 16, 18, 30, 16, 8, 10, 9]
ENC_CON = ["Hora", "Conteo", "Ubicación / nota", "Operario", "Código", "SKU", "Artículo", "Color", "Talle",
           "Marca", "Cantidad"]
ANCHO_CON = [8, 10, 22, 14, 16, 18, 30, 16, 8, 10, 9]


def _filas_desp(lista: list[dict]) -> list[list]:
    return [[f["hora"], f["tracking"], f["venta"], f["comprador"], f["producto"], f["sku"], f["cantidad"],
             f["operario"], f["origen"]] for f in lista]


def _filas_sp(envios: list[EnvioManual]) -> list[list]:
    return [[e.hora, e.plataforma, e.nombre, e.articulo, e.numero, "Sí" if e.drive_ok else "No"] for e in envios]


def _filas_dev(devs: list[Devolucion]) -> list[list]:
    return [[d.creado_at.strftime("%H:%M"), f"PD-{d.id}", d.usuario.nombre if d.usuario else "",
             i.plataforma or d.plataforma, i.venta or "", i.codigo_barras, i.sku, i.articulo, i.color, i.talle, i.estado, i.cantidad]
            for d in devs for i in d.items]


def _filas_con(conteos: list[Conteo]) -> list[list]:
    return [[c.creado_at.strftime("%H:%M"), f"#{c.id}", c.nota, c.usuario.nombre if c.usuario else "",
             i.codigo_barras, i.sku, i.articulo, i.color, i.talle, i.marca, i.cantidad]
            for c in conteos for i in c.items]


def _totales_desp(desp: list[dict], sp: list[EnvioManual]) -> list[list]:
    cnt = Counter(f["tipo"] for f in desp)
    filas = [[TIPOS_DESPACHO[t], cnt.get(t, 0)] for t in ("flex", "colecta", "tiendanube")]
    if cnt.get("cambios"):
        filas.append(["Cambios", cnt["cambios"]])
    if sp:
        filas.append(["SmartPost (carga manual)", len(sp)])
    filas.append(["TOTAL", len(desp) + len(sp)])
    return filas


def _totales_dev(devs: list[Devolucion]) -> list[list]:
    por: dict[str, Counter] = defaultdict(Counter)
    for d in devs:
        for i in d.items:
            por[i.plataforma or d.plataforma or "-"][i.estado] += i.cantidad
    filas = [[p, c.get("OPTIMO", 0), c.get("FALLA", 0), sum(c.values())] for p, c in sorted(por.items())]
    filas.append(["TOTAL", sum(f[1] for f in filas), sum(f[2] for f in filas), sum(f[3] for f in filas)])
    return filas


def _por_articulo(conteos: list[Conteo]) -> list[list]:
    acc: dict[tuple, int] = defaultdict(int)
    for c in conteos:
        for i in c.items:
            acc[(i.codigo_barras or i.sku, i.sku, i.articulo, i.color, i.talle, i.marca)] += i.cantidad
    return [list(k) + [v] for k, v in sorted(acc.items(), key=lambda kv: (kv[0][2], kv[0][3], kv[0][4]))]


def excel_despachos(db: Session, fecha: date) -> bytes:
    desp, sp = despachos_del_dia(db, fecha), smartpost_del_dia(db, fecha)
    wb = Workbook()
    _resumen(wb, "Despachos", fecha, [(["Tipo", "Despachados"], _totales_desp(desp, sp))])
    for t, nombre in TIPOS_DESPACHO.items():
        lista = [f for f in desp if f["tipo"] == t]
        if t != "cambios" or lista:  # Flex / Colecta / Tienda Nube siempre (aunque estén vacías)
            _hoja(wb, nombre, ENC_DESP, _filas_desp(lista), ANCHO_DESP)
    if sp:
        _hoja(wb, "SmartPost", ENC_SP, _filas_sp(sp), ANCHO_SP)
    return _bytes(wb)


def excel_devoluciones(db: Session, fecha: date) -> bytes:
    devs = devoluciones_del_dia(db, fecha)
    wb = Workbook()
    _resumen(wb, "Devoluciones", fecha, [(["Plataforma", "Óptimo", "Falla", "Total prendas"], _totales_dev(devs)),
                                         (["Devoluciones cargadas", len(devs)], [])])
    _hoja(wb, "Detalle", ENC_DEV, _filas_dev(devs), ANCHO_DEV)
    return _bytes(wb)


def excel_conteos(db: Session, fecha: date) -> bytes:
    conteos = conteos_del_dia(db, fecha)
    wb = Workbook()
    _resumen(wb, "Conteos", fecha, [(["Conteo", "Hora", "Ubicación / nota", "Operario", "Prendas"],
                                     [[f"#{c.id}", c.creado_at.strftime("%H:%M"), c.nota,
                                       c.usuario.nombre if c.usuario else "", c.total] for c in conteos]
                                     + [["TOTAL", "", "", "", sum(c.total for c in conteos)]])])
    _hoja(wb, "Por artículo", ["Código", "SKU", "Artículo", "Color", "Talle", "Marca", "Cantidad"],
          _por_articulo(conteos), [16, 18, 30, 16, 8, 10, 9])
    _hoja(wb, "Detalle", ENC_CON, _filas_con(conteos), ANCHO_CON)
    return _bytes(wb)


def excel_cierre_dia(db: Session, fecha: date) -> bytes:
    """Todo lo realizado y despachado en el día, en un solo archivo."""
    desp, sp = despachos_del_dia(db, fecha), smartpost_del_dia(db, fecha)
    devs, conteos = devoluciones_del_dia(db, fecha), conteos_del_dia(db, fecha)
    d = datos_cierre(db, fecha)  # tareas, actividad, pendientes de lotes
    pend = d["pendientes"]
    tareas_tot = [[d["secciones"].get(s, s), v["hechas"], v["pendientes"], v["vencidas"]]
                  for s, v in d["tareas_sec"].items()]
    ops = [[op, c.get("escaneo_ok", 0) + c.get("despacho_manual", 0), c.get("tarea_hecha", 0), sum(c.values())]
           for op, c in d["por_operario"].items()]
    prendas_dev = sum(i.cantidad for x in devs for i in x.items)
    wb = Workbook()
    _resumen(wb, "Cierre del día", fecha, [
        (["Concepto", "Cantidad"], [["Envíos despachados", len(desp) + len(sp)],
                                    ["Sin despachar (lotes)", len(pend)],
                                    ["Devoluciones", f"{len(devs)} ({prendas_dev} prendas)"],
                                    ["Conteos", f"{len(conteos)} ({sum(c.total for c in conteos)} prendas)"],
                                    ["Tareas hechas", sum(t[1] for t in tareas_tot)]]),
        (["Despacho", "Despachados"], _totales_desp(desp, sp)),
        (["Sección", "Tareas hechas", "Pendientes", "Vencidas"], tareas_tot),
        (["Operario", "Paquetes despachados", "Tareas completadas", "Acciones totales"], ops),
    ])
    _hoja(wb, "Despachados", ["Tipo"] + ENC_DESP,
          [[TIPOS_DESPACHO.get(f["tipo"], f["tipo"])] + r for f, r in zip(desp, _filas_desp(desp))], [12] + ANCHO_DESP)
    if sp:
        _hoja(wb, "SmartPost", ENC_SP, _filas_sp(sp), ANCHO_SP)
    _hoja(wb, "Sin despachar", ["Tipo", "Lote", "Nº de envío", "Venta", "Comprador", "Producto", "SKU", "Cant."],
          [[TIPOS_LOTE.get(p.lote.tipo, p.lote.tipo), p.lote.nombre, p.tracking, p.venta_id, p.comprador,
            p.producto, p.sku, p.cantidad] for p in pend], [12, 26, 20, 18, 26, 40, 18, 7])
    _hoja(wb, "Devoluciones", ENC_DEV, _filas_dev(devs), ANCHO_DEV)
    _hoja(wb, "Conteos", ENC_CON, _filas_con(conteos), ANCHO_CON)
    _hoja(wb, "Tareas", ["Sección", "Tarea", "Origen", "Estado", "Asignada a", "Vence", "Hecha por", "Hora"],
          [[d["secciones"].get(t.seccion_slug, t.seccion_slug), t.titulo, t.origen,
            "Vencida" if t.vencida(fecha) else t.estado, t.asignado_a.nombre if t.asignado_a else "",
            t.vence.strftime("%d/%m") if t.vence else "", t.hecha_por.nombre if t.hecha_por else "",
            _hhmm(t.hecha_at)] for t in d["tareas"]], [16, 50, 8, 10, 14, 8, 14, 8])
    _hoja(wb, "Actividad", ["Hora", "Operario", "Sección", "Acción", "Detalle"],
          [[a.ts.strftime("%H:%M:%S"), a.usuario.nombre if a.usuario else "", a.seccion_slug,
            ACCIONES.get(a.accion, a.accion), a.detalle] for a in d["actividad"]], [10, 14, 12, 24, 60])
    return _bytes(wb)


# ─────────────── Subida ───────────────

def _huella(db: Session, fecha: date, seccion: str) -> str:
    """Resumen de los datos de la sección (no del .xlsx, que cambia en cada armado por la hora)."""
    if seccion == "despachos":
        datos = [[f["tipo"], f["tracking"], f["sku"], f["operario"]] for f in despachos_del_dia(db, fecha)] \
            + _filas_sp(smartpost_del_dia(db, fecha))
    elif seccion == "devoluciones":
        datos = _filas_dev(devoluciones_del_dia(db, fecha))
    elif seccion == "conteos":
        datos = _filas_con(conteos_del_dia(db, fecha))
    else:  # el cierre incluye todo lo anterior + tareas y actividad
        d = datos_cierre(db, fecha)
        propias = sum(1 for a in d["actividad"] if a.accion.startswith("drive_deposito"))  # no re-subir por subir
        datos = [[_huella(db, fecha, s) for s in SECCIONES[:3]], len(d["pendientes"]), len(d["actividad"]) - propias,
                 [(t.id, t.estado) for t in d["tareas"]]]
    return hashlib.sha1(json.dumps(datos, default=str).encode("utf-8")).hexdigest()


ARMADORES = {"despachos": excel_despachos, "devoluciones": excel_devoluciones,
             "conteos": excel_conteos, "cierre": excel_cierre_dia}


def sincronizar(db: Session, fecha: date, usuario: Usuario | None = None, forzar: bool = False,
                secciones: tuple[str, ...] = SECCIONES) -> dict:
    """Arma y sube los Excel del día que cambiaron desde la última vez (o todos, con forzar=True).
    Devuelve {"fecha", "archivos": [{seccion, nombre, url}], "sin_cambios": [secciones]}."""
    if not configurado(db):
        raise DepositoDriveError("Falta conectar el Apps Script de Drive (Conexiones → Recolección Ecommerce: URL y TOKEN)")
    f = fecha.isoformat()
    if not forzar and not (despachos_del_dia(db, fecha) or smartpost_del_dia(db, fecha)
                           or devoluciones_del_dia(db, fecha) or conteos_del_dia(db, fecha)):
        return {"fecha": f, "archivos": [], "sin_cambios": list(secciones)}  # día sin movimiento: no se suben vacíos
    huellas = {s: _huella(db, fecha, s) for s in secciones}
    cambiadas = [s for s in secciones if forzar or _subido.get((f, s)) != huellas[s]]
    if not cambiadas:
        return {"fecha": f, "archivos": [], "sin_cambios": list(secciones)}
    archivos = {}
    for sec in cambiadas:
        try:
            archivos[sec] = ARMADORES[sec](db, fecha)
        except Exception as e:  # noqa: BLE001 — se dice qué archivo y por qué, en vez de un error 500 mudo
            log.exception("No se pudo armar %s del %s", sec, f)
            raise DepositoDriveError(f"No se pudo armar el Excel de {sec} del {f}: {type(e).__name__}: {e}"[:300]) from e
    try:
        r = re_drive.guardar_deposito(db, f, archivos)
    except re_drive.REDriveError as e:
        if usuario:  # el automático solo deja el aviso en el log (si no, llenaría la actividad cada 15 min)
            registrar(db, usuario, "drive_deposito_error", f"{f}: {e}"[:300], "cierre")
        raise DepositoDriveError(str(e)) from e
    for s in cambiadas:
        _subido[(f, s)] = huellas[s]
    subidos = [{"seccion": a.get("seccion"), "nombre": a.get("nombre"), "url": a.get("url")}
               for a in r.get("archivos") or []]
    registrar(db, usuario, "drive_deposito", f"{f} · " + ", ".join(a["nombre"] or "" for a in subidos), "cierre")
    return {"fecha": f, "archivos": subidos, "sin_cambios": [s for s in secciones if s not in cambiadas],
            "carpeta": r.get("carpeta", "")}


# ─────────────── Mail del cierre ───────────────

def html_cierre_dia(db: Session, fecha: date, links: dict[str, str] | None = None) -> str:
    """Cuerpo del mail: totales del día + links a los Excel de DEPOSITO PARKA (el detalle va adjunto)."""
    from html import escape
    desp, sp = despachos_del_dia(db, fecha), smartpost_del_dia(db, fecha)
    devs, conteos = devoluciones_del_dia(db, fecha), conteos_del_dia(db, fecha)
    pend = datos_cierre(db, fecha)["pendientes"]

    def tabla(enc: list[str], filas: list[list]) -> str:
        th = "".join(f"<th style='text-align:left;padding:6px 10px;background:#1f2937;color:#fff'>{escape(h)}</th>" for h in enc)
        trs = "".join("<tr>" + "".join(f"<td style='padding:6px 10px;border-bottom:1px solid #e5e7eb'>{escape(str(c))}</td>"
                                       for c in f) + "</tr>" for f in filas)
        return f"<table style='border-collapse:collapse;font-size:14px;margin:8px 0 20px'><tr>{th}</tr>{trs}</table>"

    pend_f = [[TIPOS_LOTE.get(p.lote.tipo, p.lote.tipo), p.tracking, p.comprador, (p.producto or "")[:60]] for p in pend[:50]]
    extra = f"<p style='color:#6b7280'>…y {len(pend) - 50} más (ver Excel adjunto)</p>" if len(pend) > 50 else ""
    nombres = {"despachos": "Despachos", "devoluciones": "Devoluciones", "conteos": "Conteos", "cierre": "Cierre"}
    enlaces = " · ".join(f"<a href='{escape(u)}'>{nombres.get(k, k)}</a>" for k, u in (links or {}).items() if u)
    return f"""<div style="font-family:Arial,sans-serif;color:#111">
<h2 style="margin:0 0 4px">Cierre del día — {fecha.strftime('%d/%m/%Y')}</h2>
<p style="color:#6b7280;margin:0 0 16px">PARKA Depósito · generado {ahora().strftime('%H:%M')}</p>
<h3>Despachos</h3>{tabla(['Tipo', 'Despachados'], _totales_desp(desp, sp))}
<h3>Devoluciones</h3>{tabla(['Plataforma', 'Óptimo', 'Falla', 'Total prendas'], _totales_dev(devs)) if devs else '<p>Sin devoluciones.</p>'}
<h3>Conteos</h3>{tabla(['Conteo', 'Ubicación / nota', 'Operario', 'Prendas'], [[f"#{c.id}", c.nota, c.usuario.nombre if c.usuario else "", c.total] for c in conteos]) if conteos else '<p>Sin conteos.</p>'}
<h3>Sin despachar</h3>{tabla(['Tipo', 'Nº de envío', 'Comprador', 'Producto'], pend_f) if pend_f else '<p>Todo despachado ✔</p>'}{extra}
{f"<p>Archivos en Drive (DEPOSITO PARKA): {enlaces}</p>" if enlaces else ""}
<p style="color:#6b7280">El detalle completo está en el Excel adjunto.</p></div>"""


def enviar_cierre_mail(db: Session, fecha: date, usuario: Usuario | None = None, automatico: bool = False,
                       para: str | None = None) -> "Cierre":
    """Guarda los Excel del día en Drive y manda el cierre por mail (Apps Script → MailApp) con el Excel adjunto.
    Queda registrado en la tabla de cierres (lo muestra Cierre diario → Envíos anteriores)."""
    from ..models import Cierre
    para = (para or ajustes.leer("cierre_mail_para", db)).strip()
    copia = ajustes.leer("cierre_mail_copia", db).strip()
    registro = Cierre(fecha=fecha, destinatario=para + (f" (cc {copia})" if copia else ""), automatico=automatico)
    try:
        if not configurado(db):
            raise DepositoDriveError("Falta conectar el Apps Script de Drive (Conexiones → Recolección Ecommerce)")
        links = {}
        try:  # los links del mail apuntan a los archivos recién guardados; si Drive falla igual sale el mail
            r = sincronizar(db, fecha, usuario, forzar=True)
            links = {a["seccion"]: a["url"] for a in r["archivos"]}
        except DepositoDriveError as e:
            log.warning("Cierre por mail: no se pudo guardar en Drive: %s", e)
        desp = despachos_del_dia(db, fecha) + smartpost_del_dia(db, fecha)
        devs = devoluciones_del_dia(db, fecha)
        asunto = (f"Cierre PARKA Depósito — {fecha.strftime('%d/%m/%Y')} · {len(desp)} despachados, "
                  f"{sum(i.cantidad for d in devs for i in d.items)} prendas devueltas")
        re_drive.enviar_mail(db, para, asunto, html_cierre_dia(db, fecha, links),
                             {f"{fecha.isoformat()} - Cierre.xlsx": excel_cierre_dia(db, fecha)}, copia=copia)
        registro.ok = True
        registrar(db, usuario, "cierre_mail", f"{fecha} → {registro.destinatario}"
                  + (" (automático)" if automatico else ""), "cierre", commit=False)
    except (DepositoDriveError, re_drive.REDriveError) as e:
        registro.ok, registro.error = False, str(e)[:1000]
    db.add(registro)
    db.commit()
    return registro


def cierre_mail_ya_enviado(db: Session, fecha: date) -> bool:
    """Para no mandar dos veces el automático (reinicios o más de una instancia)."""
    from ..models import Cierre
    return db.scalars(select(Cierre).where(Cierre.fecha == fecha, Cierre.automatico.is_(True),
                                           Cierre.ok.is_(True))).first() is not None
