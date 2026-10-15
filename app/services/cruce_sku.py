"""Cruce de SKU: el SKU de las publicaciones de Mercado Libre / Tienda Nube ↔ la variante de Odoo.

Los SKU viejos de ML (M-114-BLACK-XL) no son los de Odoo (1-CAMP-M-00114-BLCK-XL-PARK). Esta pantalla junta
todos los SKU que llegaron en las ventas (Recolección Ecommerce) y muestra para cada uno:

  vinculado   alguien ya lo confirmó (tabla sku_vinculos): es el que usa la Recolección
  igual       el SKU de ML ya es el de Odoo (o es un código de barras)
  automatico  la app lo encontró sola por el formato género-número-color-talle (Parka primero)
  sugerido    no tiene ese formato: se propone por el título de la publicación + color y talle (hay que mirarlo)
  sin_cruce   no se encontró nada: hay que elegirlo a mano

No escribe en Mercado Libre ni en Odoo: el cruce queda en la app.
"""
from __future__ import annotations

import io
import re
from collections import defaultdict

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import Articulo, ControlItem, ControlPaquete, SkuVinculo, Usuario
from . import mapeo
from .devoluciones import color_talle, marca

ESTADOS = ("vinculado", "igual", "automatico", "sugerido", "sin_cruce")
TIPOS_VENTA = ("flex", "colecta", "tiendanube", "ml")
_VAR_COLOR = re.compile(r"color[^:]*:\s*([^·|,]+)", re.I)
_VAR_TALLE = re.compile(r"(?:talle|size|tamaño)[^:]*:\s*([^·|,]+)", re.I)


class _Catalogo:
    """El catálogo activo en memoria, indexado para cruzar miles de SKU sin ir a la base por cada uno."""

    def __init__(self, db: Session):
        arts = db.scalars(select(Articulo).where(Articulo.inactivo.is_not(True))).all()
        self.por_id = {a.id: a for a in arts}
        self.por_sku = {a.sku.upper(): a for a in arts if a.sku}
        self.por_codigo = {a.codigo_barras: a for a in arts if a.codigo_barras}
        self.por_gen_num: dict[tuple[str, str], list[Articulo]] = defaultdict(list)
        self.por_modelo: dict[str, list[Articulo]] = defaultdict(list)
        for a in arts:
            partes = (a.sku or "").upper().split("-")
            if len(partes) >= 6:
                self.por_gen_num[(partes[2], partes[3])].append(a)
            if a.articulo:
                self.por_modelo[a.articulo.upper()].append(a)
        # nombres de modelo más largos primero: "ABBA 02003" antes que "ABBA"
        self.modelos = sorted((m for m in self.por_modelo if len(m) >= 4), key=len, reverse=True)

    def por_formato_ml(self, sku: str) -> list[Articulo]:
        """Mismo criterio que mapeo.candidatos_por_sku_ml, en memoria."""
        m = mapeo._SKU_ML.match((sku or "").strip().upper())
        if not m:
            return []
        genero, num, color = m.group(1), m.group(2).zfill(5), m.group(3)
        if genero == "B":
            genero = "U"
        talle = mapeo._talle_norm(m.group(4) or "")
        out = []
        for a in self.por_gen_num.get((genero, num), []):
            partes = a.sku.upper().split("-")
            if not talle:
                if len(partes) == 6 and mapeo.colores_equivalentes(partes[4], color):
                    out.append(a)
            elif len(partes) >= 7 and mapeo._talle_norm(partes[5]) == talle and mapeo.colores_equivalentes(partes[4], color):
                out.append(a)
        return sorted(out, key=lambda a: (mapeo._orden_marca(a), a.id))

    def por_titulo(self, titulo: str, variante: str) -> list[Articulo]:
        """El nombre del modelo aparece en el título de la publicación; color y talle salen de la variante."""
        t = f" {(titulo or '').upper()} "
        modelo = next((m for m in self.modelos if f" {m} " in t), None)
        if not modelo:
            return []
        mc, mt = _VAR_COLOR.search(variante or ""), _VAR_TALLE.search(variante or "")
        color = (mc.group(1).strip() if mc else "")
        talle = mapeo._talle_norm(mt.group(1).strip() if mt else "")
        out = []
        for a in self.por_modelo[modelo]:
            c, tl = color_talle(a)
            if talle and mapeo._talle_norm(tl) != talle:
                continue
            if color and not (mapeo.colores_equivalentes(c, color) or mapeo.colores_equivalentes(a.color, color)
                              or mapeo.colores_equivalentes(a.color_nombre, color)):
                continue
            out.append(a)
        return sorted(out, key=lambda a: (mapeo._orden_marca(a), a.id))


def art_json(a: Articulo | None) -> dict | None:
    if a is None:
        return None
    color, talle = color_talle(a)
    return {"id": a.id, "sku": a.sku, "articulo": a.articulo, "color": color, "talle": talle, "marca": marca(a),
            "codigo_barras": a.codigo_barras, "nombre": a.nombre}


def _skus_de_ventas(db: Session) -> dict[str, dict]:
    """SKU distintos que llegaron en las ventas de ML / TN, con un título y una variante de ejemplo."""
    filas = db.execute(
        select(func.upper(ControlItem.sku), func.max(ControlItem.producto), func.max(ControlItem.variante),
               func.count(), func.max(ControlPaquete.tipo))
        .join(ControlPaquete, ControlItem.paquete_id == ControlPaquete.id)
        .where(ControlItem.sku != "", ControlPaquete.tipo.in_(TIPOS_VENTA))
        .group_by(func.upper(ControlItem.sku))).all()
    return {sku.strip(): {"producto": prod or "", "variante": var or "", "ventas": n, "canal": tipo or ""}
            for sku, prod, var, n, tipo in filas if sku and sku.strip()}


def listar(db: Session) -> dict:
    cat = _Catalogo(db)
    ventas = _skus_de_ventas(db)
    vinculos = {v.sku_externo.upper(): v for v in db.scalars(select(SkuVinculo)).all()}
    usuarios = {u.id: u.nombre for u in db.scalars(select(Usuario)).all()}
    filas = []
    for sku in sorted(set(ventas) | set(vinculos)):
        info = ventas.get(sku, {"producto": "", "variante": "", "ventas": 0, "canal": ""})
        v = vinculos.get(sku)
        cands: list[Articulo] = []
        if v is not None:
            estado, art = "vinculado", cat.por_id.get(v.articulo_id)
        elif sku in cat.por_sku or (sku.isdigit() and sku in cat.por_codigo):
            estado, art = "igual", cat.por_sku.get(sku) or cat.por_codigo.get(sku)
        else:
            cands = cat.por_formato_ml(sku)
            estado = "automatico" if cands else ""
            if not cands:
                cands = cat.por_titulo(info["producto"], info["variante"])
                estado = "sugerido" if cands else "sin_cruce"
            art = cands[0] if cands else None
        filas.append({"sku": sku, **info, "estado": estado, "articulo": art_json(art),
                      "otros": [art_json(a) for a in cands[1:6]],
                      "por": usuarios.get(v.creado_por_id, "") if v else "",
                      "fecha": v.creado_at.strftime("%d/%m/%Y") if v and v.creado_at else ""})
    cuentas = {e: sum(1 for f in filas if f["estado"] == e) for e in ESTADOS}
    return {"filas": filas, "cuentas": cuentas, "total": len(filas)}


def confirmar(db: Session, sku: str, articulo_id: int, u: Usuario) -> int:
    """Guarda el cruce y lo aplica a las líneas de las recolecciones abiertas que tienen ese SKU."""
    from ..models import Control
    from .control import aplicar_articulo
    art = db.get(Articulo, articulo_id)
    if art is None:
        raise ValueError("Ese artículo no existe")
    clave = sku.strip().upper()
    if not clave:
        raise ValueError("Falta el SKU")
    mapeo.vincular_sku(db, clave, art.id, u.id)
    n = 0
    abiertos = db.scalars(select(ControlItem).join(ControlPaquete).join(Control).where(
        Control.estado == "abierto", func.upper(ControlItem.sku) == clave, ControlItem.escaneado == 0)).all()
    for it in abiertos:
        aplicar_articulo(db, it, art)
        n += 1
    db.commit()
    return n


def quitar(db: Session, sku: str) -> bool:
    v = db.scalar(select(SkuVinculo).where(SkuVinculo.sku_externo == sku.strip().upper()))
    if v is None:
        return False
    db.delete(v)
    db.commit()
    return True


def confirmar_automaticos(db: Session, u: Usuario) -> int:
    """Confirma de una vez todos los que la app cruzó sola por el formato del SKU."""
    n = 0
    for f in listar(db)["filas"]:
        if f["estado"] == "automatico" and f["articulo"]:
            confirmar(db, f["sku"], f["articulo"]["id"], u)
            n += 1
    return n


NOMBRES_ESTADO = {"vinculado": "Confirmado", "igual": "Igual a Odoo", "automatico": "Automático",
                  "sugerido": "Sugerido (revisar)", "sin_cruce": "Sin cruce"}


def excel(db: Session) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Cruce SKU"
    ws.append(["SKU ML / TN", "Estado", "SKU Odoo", "Código de barras", "Artículo", "Color", "Talle", "Marca",
               "Título de la publicación", "Variante", "Ventas"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in listar(db)["filas"]:
        a = f["articulo"] or {}
        ws.append([f["sku"], NOMBRES_ESTADO.get(f["estado"], f["estado"]), a.get("sku", ""), a.get("codigo_barras", ""),
                   a.get("articulo", ""), a.get("color", ""), a.get("talle", ""), a.get("marca", ""),
                   f["producto"], f["variante"], f["ventas"]])
    for col, ancho in zip("ABCDEFGHIJK", (24, 16, 32, 15, 20, 14, 7, 10, 50, 28, 7)):
        ws.column_dimensions[col].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def importar(db: Session, data: bytes, u: Usuario) -> dict:
    """Excel con columnas "SKU ML / TN" y "SKU Odoo" (o "Código de barras"): confirma cada fila."""
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        raise ValueError("El Excel está vacío")
    enc = [str(c or "").strip().lower() for c in filas[0]]

    def col(*claves):
        return next((i for i, c in enumerate(enc) if any(k in c for k in claves)), None)
    i_ml, i_odoo, i_cod = col("ml", "tn", "externo"), col("odoo"), col("barras", "código", "codigo")
    if i_ml is None or (i_odoo is None and i_cod is None):
        raise ValueError('Faltan las columnas "SKU ML / TN" y "SKU Odoo" (o "Código de barras")')
    cat = _Catalogo(db)
    ok, errores = 0, []
    for n, fila in enumerate(filas[1:], start=2):
        sku = str(fila[i_ml] or "").strip() if i_ml < len(fila) else ""
        if not sku:
            continue
        odoo_sku = str(fila[i_odoo] or "").strip().upper() if i_odoo is not None and i_odoo < len(fila) else ""
        cod = str(fila[i_cod] or "").strip() if i_cod is not None and i_cod < len(fila) else ""
        art = cat.por_sku.get(odoo_sku) or cat.por_codigo.get(cod)
        if art is None:
            if odoo_sku or cod:
                errores.append(f"Fila {n}: {odoo_sku or cod} no está en el catálogo")
            continue
        confirmar(db, sku, art.id, u)
        ok += 1
    return {"confirmados": ok, "errores": errores[:50]}
